#!/usr/bin/env python3
"""One-off: pull 1_Companies and Events_Master out of the .xlsm into silver CSVs.

python migrate_from_excel.py path/to/20260724_securitymaster_v03.xlsm
"""

import csv
import sys

from openpyxl import load_workbook

from biointel import config
from biointel.pipeline import COMPANY_COLS, EVENT_COLS


def main(path):
    wb = load_workbook(path, data_only=True)

    co = wb["1_Companies"]
    companies = []
    for r in range(2, co.max_row + 1):
        iid = co.cell(row=r, column=1).value
        if iid is None:
            continue
        created = co.cell(row=r, column=4).value
        companies.append(
            {
                "IID": iid,
                "Name": co.cell(row=r, column=2).value,
                "Ticker": co.cell(row=r, column=3).value,
                "Created": created.date().isoformat() if hasattr(created, "date") else created,
                "Description": co.cell(row=r, column=5).value,
                "CIK": "",
                "SIC": "",
                "SICDescription": "",
                "Exchange": "",
                "StateOfIncorporation": "",
            }
        )

    em = wb["Events_Master"]
    events = []
    for r in range(2, em.max_row + 1):
        iid = em.cell(row=r, column=1).value
        if iid is None:
            continue
        d = em.cell(row=r, column=4).value
        events.append(
            {
                "IID": iid,
                "Name": em.cell(row=r, column=2).value,
                "Event": em.cell(row=r, column=3).value,
                "Date": d.date().isoformat() if hasattr(d, "date") else d,
                "AppNo": em.cell(row=r, column=5).value,
                "Drug": em.cell(row=r, column=6).value,
                "Outcome": em.cell(row=r, column=7).value,
                "Priority": em.cell(row=r, column=8).value,
                "ClassCode": em.cell(row=r, column=9).value,
                "SubType": em.cell(row=r, column=10).value,
            }
        )

    for path_, cols, rows in (
        (config.COMPANIES_CSV, COMPANY_COLS, companies),
        (config.EVENTS_CSV, EVENT_COLS, events),
    ):
        with path_.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
        print(f"{len(rows):>5} rows -> {path_}")

    print(
        "\nNote: CIK/SIC/Exchange are blank for migrated rows. They fill in\n"
        "for companies added via `python cli.py add`."
    )


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1])
